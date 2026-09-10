"""Genera plantillas .docx de arranque en plantillas/.

Son borradores funcionales con la sintaxis Jinja que espera `services/documentos.py`;
el abogado las reemplaza por las minutas reales conservando las variables.

    python scripts/crear_plantillas.py [--forzar]
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document  # noqa: E402

from app.config import settings  # noqa: E402

DIAGNOSTICO = [
    ("DIAGNOSTICO LEGAL PRELIMINAR", "Title"),
    ("HENAO LEGAL - Industria musical", "Subtitle"),
    ("Talento: {{ talento_nombre }}", None),
    ("Documento: {{ talento_documento }}", None),
    ("Fecha de intake: {{ fecha_intake }}", None),
    ("Fecha del diagnostico: {{ fecha_larga }}", None),
    ("Resumen ejecutivo", "Heading 1"),
    ("{{ resumen_ejecutivo }}", None),
    ("Resumen por area", "Heading 1"),
    ("{% for fila in resumen_estados %}{{ fila.area }} — estado {{ fila.estado }} — {{ fila.pendientes }} pendientes — prioridad {{ fila.prioridad }}", None),
    ("{% endfor %}", None),
    ("Analisis por area", "Heading 1"),
    ("{% for area in areas %}", None),
    ("{{ area.titulo }} (estado global: {{ area.estado_global }})", "Heading 2"),
    ("{{ area.observaciones }}", None),
    ("{% for a in area.aspectos %}- [{{ a.estado }}] {{ a.aspecto }}{% if a.prioridad %} ({{ a.prioridad }}, {{ a.plazo_sugerido }}){% endif %}: {{ a.accion_requerida or 'sin accion pendiente' }} — responsable: {{ a.responsable }}", None),
    ("{% endfor %}{% endfor %}", None),
    ("Plan de accion consolidado", "Heading 1"),
    ("{% for item in plan_de_accion %}{{ loop.index }}. [{{ item.prioridad }}] {{ item.aspecto }} ({{ item.area }}) — {{ item.accion_requerida }} — {{ item.responsable }}, {{ item.plazo_sugerido }}", None),
    ("{% endfor %}", None),
    (
        "Este documento es un diagnostico preliminar sujeto a revision del abogado "
        "responsable. No constituye un concepto juridico definitivo.",
        None,
    ),
]

CONTRATO = [
    ("CONTRATO DE PRESTACION DE SERVICIOS JURIDICOS", "Title"),
    (
        "Entre HENAO LEGAL, en adelante EL DESPACHO, y {{ talento_nombre }}, "
        "identificado con documento {{ talento_documento }}, en adelante EL TALENTO, "
        "se celebra el presente contrato el {{ fecha_larga }}.",
        None,
    ),
    ("PRIMERA. Objeto", "Heading 1"),
    (
        "EL DESPACHO prestara a EL TALENTO servicios juridicos en materia de industria "
        "musical, conforme al diagnostico legal preliminar del {{ fecha_hoy }}.",
        None,
    ),
    ("SEGUNDA. Alcance", "Heading 1"),
    ("{% for item in plan_de_accion %}- {{ item.aspecto }} ({{ item.area }})", None),
    ("{% endfor %}", None),
    ("TERCERA. Honorarios", "Heading 1"),
    ("[Completar por el despacho]", None),
    ("Firmas", "Heading 1"),
    ("EL TALENTO: {{ talento_nombre }}     EL DESPACHO: HENAO LEGAL", None),
]

PODER = [
    ("PODER ESPECIAL", "Title"),
    (
        "Yo, {{ talento_nombre }}, identificado con documento {{ talento_documento }}, "
        "confiero poder especial, amplio y suficiente a HENAO LEGAL para que me "
        "represente en los asuntos derivados de mi actividad en la industria musical, "
        "en particular:",
        None,
    ),
    ("{% for item in plan_de_accion %}- {{ item.aspecto }}", None),
    ("{% endfor %}", None),
    ("Dado el {{ fecha_larga }}.", None),
    ("_______________________________", None),
    ("{{ talento_nombre }}", None),
]

PLANTILLAS = {
    "diagnostico.docx": DIAGNOSTICO,
    "contrato.docx": CONTRATO,
    "poder.docx": PODER,
}


def construir(bloques: list[tuple[str, str | None]], destino: Path) -> None:
    doc = Document()
    for texto, estilo in bloques:
        if estilo:
            doc.add_paragraph(texto, style=estilo)
        else:
            doc.add_paragraph(texto)
    doc.save(str(destino))


def main() -> None:
    forzar = "--forzar" in sys.argv
    settings.plantillas_dir.mkdir(parents=True, exist_ok=True)
    for nombre, bloques in PLANTILLAS.items():
        destino = settings.plantillas_dir / nombre
        if destino.exists() and not forzar:
            print(f"- {nombre}: ya existe, se conserva (usa --forzar para sobrescribir)")
            continue
        construir(bloques, destino)
        print(f"- {nombre}: generada en {destino}")


if __name__ == "__main__":
    main()
